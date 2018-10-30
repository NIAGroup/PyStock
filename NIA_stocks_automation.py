################################################################################
####remember to include the "s" in the url so the proxy works effectively.
####REMEMBER TO INSTALL the alpha_vantage library for this code.
####alpha_vantage is a library we've found that's free and collects stock data.
################################################################################

import json, requests, alpha_vantage, openpyxl, argparse
import matplotlib.pyplot as plt
import matplotlib
from matplotlib import dates
from openpyxl import Workbook


def getArgs():
    parser = argparse.ArgumentParser(description='stuff')
    parser.add_argument('-s')
    args = parser.parse_args()
    s = args.s
    return s

def getLeTicks(array_corrected_x):
	leTicks = []
	for t in range(0,len(array_corrected_x)):
		if t%4 == 0:
			leTicks.append(array_corrected_x[t])

	return leTicks

stock_symbol = getArgs()
row = 1
col = 1
wb = Workbook()
ws = wb.create_sheet("Main")
share_array = []
year_array = []
array_corrected_y = []
array_corrected_x = []
API_URL = "https://www.alphavantage.co/query"

data = {
	"function": "SMA",
	"time_period": 10,
	"series_type": "open",
	"symbol": stock_symbol,
	"interval": "monthly",
	"apikey": "ksjendcksnfk",
}

proxies = {
	"https" : "http://proxy-us.intel.com:911"
}
try:
    response = requests.get(API_URL, params=data,proxies=proxies)
except Exception:
    response = requests.get(API_URL, params=data)
arr = response.text.split("{")

js = json.loads(response.text)
d = js["Technical Analysis: SMA"]
row = 1
col = 1

for s in d:
	if(row>1):
	    stringer_a = str(s)
	    stringer_b = str(d[s]["SMA"])
	    year_array.append(str(stringer_a))
	    share_array.append(float(stringer_b))
	    print(stringer_a+"\t"+stringer_b)
	    ws.cell(row-1,col,stringer_a)
	    ws.cell(row-1,col+1,stringer_b)
	row+=1
wb.save("results.xlsx")
year_array =  matplotlib.dates.datestr2num(year_array)
formatter = dates.DateFormatter('%Y-%m-%d')
for x in range(len(year_array)-1,-1,-1):
	array_corrected_x.append(year_array[x])
for x in range(len(share_array)-1,-1,-1):
	array_corrected_y.append(share_array[x])

array_corrected_x = getLeTicks(array_corrected_x)
array_corrected_y = getLeTicks(array_corrected_y)
plt.figure().patch.set_facecolor('xkcd:mint green')
plt.plot(array_corrected_x,array_corrected_y, color='darkorange',marker='o')
plt.xlabel("Dates")
plt.ylabel("Share Value")
plt.title("Share Value Trends for: "+stock_symbol)
plt.grid()
the_ticks = array_corrected_x
plt.xticks(the_ticks)
plt.xticks(rotation=70)
ax = plt.gcf().axes[0] 
ax.xaxis.set_major_formatter(formatter)

plt.show()
